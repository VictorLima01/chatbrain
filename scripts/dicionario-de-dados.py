"""Gera o dicionario de dados de uma pasta de material do desafio.

Os .csv e .xlsx nao sao lidos pela ingestao (o extrator suporta vtt, pdf, pptx,
md e txt). Este script le essas fontes e escreve um Markdown -- que a ingestao
le -- com a estrutura de cada tabela: colunas, contagem de linhas e algumas
linhas de exemplo. A ideia e tornar os dados pesquisaveis sem jogar 100 mil
linhas de CSV na base vetorial, o que geraria milhares de chunks sem valor.

Uso:
    python scripts/dicionario-de-dados.py <pasta-do-material>
    python scripts/dicionario-de-dados.py <pasta> --validar http://localhost:8081

O limite de 512 tokens do modelo de embeddings e o que dita o formato daqui.
Texto tabular tokeniza muito mais denso que prosa -- codigos como `A50110000`
viram varios tokens cada um -- e um chunk de 1200 caracteres de amostra bruta
passa de 600 tokens, que o TEI recusa com 413 e o documento inteiro falha. Por
isso a amostra e curta: poucos campos por linha, valor truncado. O modo
--validar passa uma janela deslizante do tamanho do chunk pelo /tokenize do TEI
e aponta a pior janela, para conferir que o arquivo gerado cabe.
"""
import csv
import glob
import json
import os
import sys
import urllib.request

AMOSTRAS = 2        # linhas de exemplo por tabela
MAX_CAMPOS = 6      # campos mostrados por linha de exemplo
MAX_VALOR = 36      # caracteres por valor
CHUNK = 1200        # igual ao rag.chunk-size do application.yml
LIMITE_TOKENS = 512  # limite do intfloat/multilingual-e5-small
PASSO_JANELA = 100  # resolucao da varredura do --validar


def ler_csv(caminho):
    """Devolve (colunas, linhas_de_amostra, total_de_linhas)."""
    for encoding in ("utf-8-sig", "latin-1"):
        try:
            with open(caminho, newline="", encoding=encoding) as f:
                amostra = f.read(8192)
                f.seek(0)
                try:
                    dialeto = csv.Sniffer().sniff(amostra, delimiters=",;\t|")
                except csv.Error:
                    dialeto = csv.excel
                leitor = csv.reader(f, dialeto)
                colunas = next(leitor, [])
                exemplos, total = [], 0
                for linha in leitor:
                    total += 1
                    if len(exemplos) < AMOSTRAS and any(c.strip() for c in linha):
                        exemplos.append(linha)
                return colunas, exemplos, total
        except UnicodeDecodeError:
            continue
    return [], [], 0


def bloco_amostra(colunas, linha):
    """Poucos campos, valores curtos: o chunk precisa caber em 512 tokens."""
    partes = []
    for i, coluna in enumerate(colunas):
        if len(partes) >= MAX_CAMPOS:
            partes.append("...")
            break
        valor = linha[i].strip() if i < len(linha) else ""
        if valor:
            if len(valor) > MAX_VALOR:
                valor = valor[:MAX_VALOR].rstrip() + "..."
            partes.append(f"{coluna.strip()} = {valor}")
    return "; ".join(partes)


def milhar(n):
    return f"{n:,}".replace(",", ".")


def gerar(base):
    """Devolve {caminho_relativo: conteudo} — um arquivo por tabela, mais o indice.

    Um arquivo por tabela nao e capricho de organizacao: e o que mantem cada
    chunk abaixo dos 512 tokens. Num arquivo unico, o chunker empilharia a
    lista de colunas de uma tabela com a da seguinte, e 1200 caracteres de
    nomes de coluna passam do limite do modelo.
    """
    arquivos = {}
    csvs = sorted(glob.glob(os.path.join(base, "04_CSVs_para_DataDrop", "*.csv")))
    resumo = []

    for caminho in csvs:
        nome_arquivo = os.path.basename(caminho)
        tabela = nome_arquivo.rsplit(".", 1)[0]
        colunas, exemplos, total = ler_csv(caminho)
        nomes = [c.strip() for c in colunas if c.strip()]
        linhas = [
            f"# {tabela} — tabela do desafio Burgubras",
            "",
            f"Fonte: `04_CSVs_para_DataDrop/{nome_arquivo}`, carregada no Apptio Costing pelo",
            f"DataDrop. {milhar(total)} linhas de dados, {len(nomes)} colunas.",
            "",
            "## Colunas",
            "",
            ", ".join(nomes) + ".",
            "",
        ]
        if exemplos:
            linhas += ["## Linhas de exemplo", ""]
            for i, exemplo in enumerate(exemplos, 1):
                linhas += [f"{i}. {bloco_amostra(colunas, exemplo)}", ""]
        arquivos[f"dicionario/{tabela}.md"] = "\n".join(linhas) + "\n"
        resumo.append((nome_arquivo, total, len(nomes)))

    # Nome do CSV equivalente a cada aba: "CMDB CI App" <-> CMDB_CI_App.csv.
    por_csv = {os.path.basename(c).rsplit(".", 1)[0].replace("_", " ").lower(): os.path.basename(c)
               for c in csvs}
    abas = []

    xlsx = sorted(glob.glob(os.path.join(base, "03_Dados_Excel", "*.xlsx")))
    if xlsx:
        import openpyxl

        for caminho in xlsx:
            wb = openpyxl.load_workbook(caminho, read_only=True, data_only=True)
            for aba in wb.worksheets:
                total = max((aba.max_row or 1) - 1, 0)
                equivalente = por_csv.get(aba.title.strip().lower())
                cabecalho = []
                if not equivalente:
                    for linha in aba.iter_rows(min_row=1, max_row=1, values_only=True):
                        cabecalho = [str(v).strip() for v in linha if v is not None]
                abas.append((os.path.basename(caminho), aba.title, total, equivalente, cabecalho))
            wb.close()

    indice = [
        "# Dicionario de dados do desafio Burgubras",
        "",
        "As fontes de dados do desafio e onde achar a estrutura de cada uma. Os `.csv` e o `.xlsx`",
        "ficam no bucket para download, mas a ingestao nao le planilha: a estrutura de cada tabela",
        "esta num arquivo proprio em `dicionario/`, gerado por `scripts/dicionario-de-dados.py`.",
        "",
        "## Tabelas em CSV (carga via DataDrop)",
        "",
    ]
    for nome_arquivo, total, n_colunas in resumo:
        tabela = nome_arquivo.rsplit(".", 1)[0]
        indice.append(f"- `{nome_arquivo}` — {milhar(total)} linhas, {n_colunas} colunas "
                      f"(estrutura em `dicionario/{tabela}.md`)")
    if abas:
        indice += [
            "",
            "## Planilha de mock data",
            "",
            "As abas sao a origem dos CSVs. Onde a aba corresponde a um CSV, as colunas sao as",
            "mesmas e nao estao repetidas: repetir o cabecalho geraria dois trechos quase identicos",
            "na base, competindo entre si na busca.",
            "",
        ]
        for planilha, aba, total, equivalente, cabecalho in abas:
            if equivalente:
                indice.append(f"- Aba `{aba}` — {milhar(total)} linhas; mesmas colunas de "
                              f"`{equivalente}`")
            else:
                indice.append(f"- Aba `{aba}` — {milhar(total)} linhas; colunas: "
                              + (", ".join(cabecalho) if cabecalho else "(vazia)"))

    arquivos["DICIONARIO_DE_DADOS.md"] = "\n".join(indice) + "\n"
    return arquivos


def validar(texto, url):
    """Passa uma janela do tamanho do chunk pelo /tokenize e devolve o pior caso."""
    pior, onde = 0, 0
    for inicio in range(0, max(len(texto) - CHUNK, 0) + 1, PASSO_JANELA):
        janela = texto[inicio:inicio + CHUNK]
        corpo = json.dumps({"inputs": janela}).encode()
        req = urllib.request.Request(
            url.rstrip("/") + "/tokenize", data=corpo,
            headers={"Content-Type": "application/json"})
        with urllib.request.urlopen(req, timeout=60) as resp:
            tokens = len(json.load(resp)[0])
        if tokens > pior:
            pior, onde = tokens, inicio
    return pior, onde


if __name__ == "__main__":
    base = sys.argv[1]
    arquivos = gerar(base)

    for relativo, conteudo in sorted(arquivos.items()):
        destino = os.path.join(base, relativo)
        os.makedirs(os.path.dirname(destino), exist_ok=True)
        with open(destino, "w", encoding="utf-8") as f:
            f.write(conteudo)
    print(f"{len(arquivos)} arquivos escritos em {base}")

    if "--validar" in sys.argv:
        url = sys.argv[sys.argv.index("--validar") + 1]
        pior_geral, pior_arquivo = 0, ""
        for relativo, conteudo in sorted(arquivos.items()):
            tokens, _ = validar(conteudo, url)
            if tokens > pior_geral:
                pior_geral, pior_arquivo = tokens, relativo
        folga = LIMITE_TOKENS - pior_geral
        print(f"pior janela de {CHUNK} chars: {pior_geral} tokens em {pior_arquivo}; "
              f"folga de {folga} para o limite de {LIMITE_TOKENS}")
        sys.exit(0 if pior_geral < LIMITE_TOKENS else 1)
